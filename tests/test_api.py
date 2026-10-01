from __future__ import annotations

import io
import logging

import pytest
from fastapi.testclient import TestClient

from bestbill.api.catalog_source import CatalogProvider, CatalogSettings
from bestbill.api.main import create_app
from bestbill.api.settings import Settings

from .api_helpers import compare_body, xlsx_bytes


def make_client(catalog_dir, **kwargs) -> TestClient:
    settings = Settings(
        catalog=CatalogSettings(catalog_path=catalog_dir / "catalog.sqlite"),
        **kwargs,
    )
    return TestClient(create_app(settings))


@pytest.fixture()
def client(catalog_dir):
    with make_client(catalog_dir) as c:
        yield c


@pytest.fixture()
def empty_client(tmp_path):
    settings = Settings(
        catalog=CatalogSettings(catalog_path=tmp_path / "missing.sqlite")
    )
    with TestClient(create_app(settings)) as c:
        yield c


# -- health / meta -----------------------------------------------------------
def test_health_ok(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok"
    assert body["catalog_loaded"] is True
    assert body["snapshot_date"]
    assert body["catalog_age_days"] >= 0


def test_health_supports_head(client):
    response = client.head("/api/health")
    assert response.status_code == 200
    assert response.content == b""


def test_health_without_catalog(empty_client):
    body = empty_client.get("/api/health").json()
    assert body == {
        "status": "ok",
        "catalog_loaded": False,
        "snapshot_date": None,
        "catalog_age_days": None,
    }


def test_catalog_meta(client):
    body = client.get("/api/catalog/meta").json()
    assert body["stale"] is False
    assert body["counts"]["included"] > 0
    assert set(body["counts"]["excluded_by_reason"])
    assert body["licence"] == "CC-BY-SA-4.0"
    assert body["sources"] and body["sources"][0]["name"]
    assert "ARERA" in body["attribution"]


def test_catalog_meta_stale_when_old(catalog_dir, monkeypatch):
    import bestbill.api.catalog_source as cs

    class FakeDate(cs.date):
        @classmethod
        def today(cls):
            return cls(2030, 1, 1)

    monkeypatch.setattr(cs, "date", FakeDate)
    with make_client(catalog_dir) as c:
        meta = c.get("/api/catalog/meta").json()
    assert meta["stale"] is True and meta["age_days"] > 3


@pytest.mark.parametrize("path", ["/api/catalog/meta", "/api/offers"])
def test_get_endpoints_503_without_catalog(empty_client, path):
    r = empty_client.get(path)
    assert r.status_code == 503
    assert "catalogo" in r.json()["detail"].lower()


def test_compare_503_without_catalog(empty_client):
    r = empty_client.post("/api/compare", json=compare_body())
    assert r.status_code == 503


# -- offers ------------------------------------------------------------------
def test_offers_pagination_and_fields(client):
    page = client.get("/api/offers", params={"limit": 5}).json()
    assert page["total"] > 5 and len(page["items"]) == 5
    assert set(page["items"][0]) == {
        "id",
        "supplier",
        "name",
        "url",
        "price_type",
        "band_structure",
        "source",
        "valid_to",
        "duration_months",
        "duration_open_ended",
    }
    page2 = client.get("/api/offers", params={"limit": 5, "offset": 5}).json()
    assert page2["items"][0]["id"] != page["items"][0]["id"]


def test_offers_filters(client):
    fixed = client.get(
        "/api/offers", params={"price_type": "fixed", "limit": 200}
    ).json()
    assert fixed["items"] and all(i["price_type"] == "fixed" for i in fixed["items"])
    placet = client.get("/api/offers", params={"source": "placet", "limit": 200}).json()
    assert all(i["source"] == "placet" for i in placet["items"])
    first = placet["items"][0]
    by_supplier = client.get(
        "/api/offers", params={"supplier": first["supplier"][:4].lower()}
    ).json()
    assert by_supplier["total"] >= 1
    by_name = client.get("/api/offers", params={"q": first["name"].upper()}).json()
    assert any(i["id"] == first["id"] for i in by_name["items"])


@pytest.mark.parametrize(
    "params", [{"limit": 0}, {"limit": 1000}, {"offset": -1}, {"price_type": "x"}]
)
def test_offers_validation(client, params):
    r = client.get("/api/offers", params=params)
    assert r.status_code == 422
    assert r.json()["detail"][0]["message"]


# -- comuni / sample ---------------------------------------------------------
def test_comuni_search(client):
    r = client.get("/api/comuni", params={"q": "GENOV"}).json()
    assert r[0] == {
        "codice": "010025",
        "nome": "Genova",
        "sigla_provincia": "GE",
        "regione": "Liguria",
    }
    assert len(client.get("/api/comuni", params={"q": "sa"}).json()) == 20


def test_comuni_requires_query(client):
    assert client.get("/api/comuni").status_code == 422
    assert client.get("/api/comuni", params={"q": "a"}).status_code == 422


def test_sample(client):
    body = client.get("/api/sample").json()
    assert body["location"] == "Esempio"
    assert len(body["months"]) == 12
    assert all(m["kwh"] > 0 and set(m) == {"month", "kwh"} for m in body["months"])


def test_sample_round_trips_into_compare(client):
    """The UI sends /api/sample (or /api/parse) months straight to /api/compare."""
    months = client.get("/api/sample").json()["months"]
    r = client.post("/api/compare", json={"consumption": months, "top_n": 3})
    assert r.status_code == 200, r.text
    assert r.json()["results"]


# -- parse -------------------------------------------------------------------
def upload(client, content: bytes, name: str = "storico.xlsx"):
    return client.post("/api/parse", files={"file": (name, io.BytesIO(content))})


def test_parse_ok(client):
    r = upload(client, xlsx_bytes("Milano", kwh=300))
    assert r.status_code == 200
    body = r.json()
    assert body["skipped"] == []
    prof = body["profiles"][0]
    assert prof["location"] == "Milano"
    assert prof["months"][0] == {"month": "2025-01", "kwh": 300.0}
    assert len(prof["months"]) == 12


def test_parse_sample_file(client):
    from pathlib import Path

    import bestbill

    path = Path(bestbill.__file__).parent / "data" / "sample.xlsx"
    r = upload(client, path.read_bytes())
    assert r.status_code == 200
    assert r.json()["profiles"][0]["location"] == "Esempio"


def test_parse_skips_unusable_location(client):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes("Ok")))
    bad = wb.create_sheet("Storico_Corta")
    bad.append(
        ["Mese", "Consumo [kWh]", "PUN mensile [€/kWh]", "FC PUN mensile [€/kWh]"]
    )
    buf = io.BytesIO()
    wb.save(buf)
    body = upload(client, buf.getvalue()).json()
    assert body["skipped"] == ["Corta"]
    assert [p["location"] for p in body["profiles"]] == ["Ok"]


def test_parse_rejects_not_xlsx(client):
    r = upload(client, b"hello, this is not a spreadsheet")
    assert r.status_code == 422
    assert r.json()["detail"][0]["field"] == "file"


def test_parse_rejects_zip_that_is_not_a_workbook(client):
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("a.txt", "x")
    assert upload(client, buf.getvalue()).status_code == 422


def test_parse_rejects_wrong_format(client):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(client, buf.getvalue())
    assert r.status_code == 422
    assert "Storico_" in r.json()["detail"][0]["message"]


def test_parse_wrong_month_count_is_422(client):
    r = upload(client, xlsx_bytes(months=5))
    assert r.status_code == 422


def test_parse_error_does_not_echo_values(client):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes()))
    ws = wb.active
    ws["B2"] = "987654.321 kWh"
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(client, buf.getvalue())
    assert r.status_code == 422
    assert "987654" not in r.text


def test_parse_missing_file(client):
    r = client.post("/api/parse")
    assert r.status_code == 422


def test_parse_upload_too_large(catalog_dir):
    with make_client(catalog_dir, max_upload_bytes=1000) as c:
        r = upload(c, xlsx_bytes())
    assert r.status_code == 413


def test_parse_file_over_cap_inside_envelope(catalog_dir):
    # Under the body-limit (cap + framing slack) but over the file cap.
    with make_client(catalog_dir, max_upload_bytes=1000) as c:
        r = upload(c, b"PK\x03\x04" + b"0" * 1500)
    assert r.status_code == 413
    assert "2 MB" in r.json()["detail"]


# -- compare -----------------------------------------------------------------
def test_compare_happy_path(client):
    r = client.post("/api/compare", json=compare_body())
    assert r.status_code == 200
    body = r.json()
    assert body["results"]
    ranks = [x["rank"] for x in body["results"]]
    assert ranks == sorted(ranks) and ranks[0] == 1
    costs = [x["cost_eur"] for x in body["results"]]
    assert costs == sorted(costs)
    top = body["results"][0]
    assert top["delta_vs_best_eur"] == 0
    assert set(top["breakdown"]) >= {"energy", "fixed_fees", "total"}
    assert isinstance(top["energy_price_eur_kwh"], float)
    assert top["energy_price_kind"] in {"fixed", "pun_spread"}
    assert top["supplier_name_source"] in {"arera", "placet", "domain", "vat"}
    a = body["assumptions"]
    assert a["cost_label"] == "costo materia energia (IVA esclusa)"
    assert a["period_start"] == "2025-01-01" and a["period_end"] == "2025-12-01"
    assert a["band_split_source"] == "standard"
    assert a["scenario"] == {"kind": "historical"}
    assert a["statement"]
    assert len(a["pun_months_used"]) == 12
    assert body["snapshot_date"]
    assert body["total_eligible"] >= body["total_matching"] == body["total_eligible"]
    assert body["excluded_count"] == sum(body["excluded_by_reason"].values())


def test_compare_top_n_and_filters(client):
    body = client.post("/api/compare", json=compare_body(top_n=3)).json()
    assert len(body["results"]) == 3
    assert body["total_matching"] > 3
    fixed = client.post(
        "/api/compare", json=compare_body(filters={"price_type": "fixed"})
    ).json()
    assert fixed["results"]
    assert all(x["price_type"] == "fixed" for x in fixed["results"])
    assert fixed["total_matching"] < fixed["total_eligible"]
    placet = client.post(
        "/api/compare", json=compare_body(filters={"source": "placet"})
    ).json()
    assert placet["results"]
    assert all(x["source"] == "placet" for x in placet["results"])


def test_compare_scenarios(client):
    flat = client.post(
        "/api/compare",
        json=compare_body(scenario={"kind": "flat", "value": 0.2}),
    ).json()
    assert set(flat["assumptions"]["pun_months_used"].values()) == {0.2}
    scaled = client.post(
        "/api/compare",
        json=compare_body(scenario={"kind": "scaled", "factor": 1.5}),
    ).json()
    assert scaled["assumptions"]["scenario"] == {"kind": "scaled", "factor": 1.5}


def test_compare_with_bands_and_residency_and_comune(client):
    consumption = [
        {"month": f"2025-{m:02d}", "kwh": 300, "f1": 100, "f2": 100, "f3": 100}
        for m in range(1, 13)
    ]
    r = client.post(
        "/api/compare",
        json=compare_body(
            consumption=consumption,
            residency="non_resident",
            istat_comune="010025",
            committed_power_kw=4.5,
        ),
    )
    assert r.status_code == 200
    a = r.json()["assumptions"]
    assert a["band_split_source"] == "user"
    assert a["residency"] == "non_resident"
    assert a["istat_comune"] == "010025"
    assert a["committed_power_kw"] == 4.5


def test_compare_comune_resolves_area_offers(client):
    # Independent of today's date: a comune turns "zona non specificata"
    # exclusions into real area checks.
    without = client.post("/api/compare", json=compare_body()).json()
    with_comune = client.post(
        "/api/compare", json=compare_body(istat_comune="010025")
    ).json()
    assert without["excluded_by_reason"]["zona non specificata"] > 0
    assert "zona non specificata" not in with_comune["excluded_by_reason"]
    assert (
        with_comune["excluded_by_reason"].get(
            "offerta non disponibile nel comune indicato", 0
        )
        > 0
    )


def test_compare_unknown_comune(client):
    r = client.post("/api/compare", json=compare_body(istat_comune="999999"))
    assert r.status_code == 422
    assert r.json()["detail"][0]["field"] == "istat_comune"


def bad(**overrides):
    return compare_body(**overrides)


def months(n=12, start=1, year=2025):
    return [{"month": f"{year}-{m:02d}", "kwh": 100} for m in range(start, start + n)]


@pytest.mark.parametrize(
    ("body", "field", "text"),
    [
        (bad(consumption=months(11)), "consumption", "12 mesi"),
        (bad(consumption=months(12)[:11] + [months(12)[0]]), "consumption", "diversi"),
        (
            bad(
                consumption=months(6)
                + months(5, start=8)
                + [{"month": "2026-01", "kwh": 1}]
            ),
            "consumption",
            "consecutivi",
        ),
        (
            bad(consumption=[{"month": "2025-13", "kwh": 1}] * 12),
            "consumption[0].month",
            "AAAA-MM",
        ),
        (
            bad(consumption=[{"month": "2025-01", "kwh": -5}] + months(12)[1:]),
            "consumption[0].kwh",
            "maggiore o uguale a 0",
        ),
        (
            bad(consumption=[{"month": "2025-01", "kwh": "abc"}] + months(12)[1:]),
            "consumption[0].kwh",
            "numero",
        ),
        (
            bad(
                consumption=[{"month": "2025-01", "kwh": 100, "f1": 100}]
                + months(12)[1:]
            ),
            "consumption[0]",
            "F1, F2 e F3",
        ),
        (
            bad(
                consumption=[
                    {"month": f"2025-{m:02d}", "kwh": 100, "f1": 10, "f2": 10, "f3": 10}
                    for m in range(1, 13)
                ]
            ),
            "consumption",
            "0,5%",
        ),
        (bad(residency="tenant"), "residency", "resident"),
        (bad(istat_comune="12"), "istat_comune", "6 cifre"),
        (bad(committed_power_kw=0), "committed_power_kw", "maggiore di 0"),
        (bad(top_n=500), "top_n", "minore o uguale a 200"),
        (bad(scenario={"kind": "wild"}), "scenario", "Scenario"),
        (
            bad(scenario={"kind": "scaled", "factor": 0}),
            "scenario.scaled.factor",
            "maggiore di 0",
        ),
        (bad(filters={"price_type": "x"}), "filters.price_type", "fixed"),
        (bad(nope=1), "nope", "non riconosciuto"),
        ({}, "consumption", "obbligatorio"),
    ],
)
def test_compare_validation_errors_are_italian(client, body, field, text):
    r = client.post("/api/compare", json=body)
    assert r.status_code == 422
    details = r.json()["detail"]
    assert any(d["field"] == field and text in d["message"] for d in details), details


def test_compare_mixed_band_months(client):
    consumption = months(12)
    consumption[0] = {"month": "2025-01", "kwh": 90, "f1": 30, "f2": 30, "f3": 30}
    r = client.post("/api/compare", json=bad(consumption=consumption))
    assert r.status_code == 422
    assert "tutti i mesi" in r.json()["detail"][0]["message"]


def test_compare_invalid_json(client):
    r = client.post(
        "/api/compare", content=b"{oops", headers={"content-type": "application/json"}
    )
    assert r.status_code == 422
    assert r.json()["detail"][0]["message"] == "JSON non valido."


def test_validation_errors_do_not_echo_input(client):
    body = bad(consumption=[{"month": "2025-01", "kwh": -123456.789}] + months(12)[1:])
    r = client.post("/api/compare", json=body)
    assert "123456" not in r.text


def test_unhandled_error_is_generic_500(client, monkeypatch, caplog):
    import bestbill.api.routes as routes

    def boom(*a, **k):
        raise RuntimeError("secret 271.5 internals")

    monkeypatch.setattr(routes, "compare", boom)
    with caplog.at_level(logging.DEBUG):
        r = client.post("/api/compare", json=compare_body())
    assert r.status_code == 500
    assert r.json() == {"detail": "Errore interno. Riprova più tardi."}
    assert "secret" not in r.text and "secret" not in caplog.text
    assert "RuntimeError" in caplog.text
    assert r.headers["x-content-type-options"] == "nosniff"


# -- privacy -----------------------------------------------------------------
def test_compare_does_not_log_consumption(client, caplog, capsys):
    body = compare_body(
        consumption=[
            {"month": f"2025-{m:02d}", "kwh": 4242.31 + m * 0.137} for m in range(1, 13)
        ]
    )
    with caplog.at_level(logging.DEBUG):
        r = client.post("/api/compare", json=body)
        client.post("/api/compare", json={**body, "top_n": 0})  # a 422 too
    assert r.status_code == 200
    out = capsys.readouterr()
    text = caplog.text + out.out + out.err
    for item in body["consumption"]:
        assert str(item["kwh"]) not in text
    assert "4242" not in text


# -- hardening ---------------------------------------------------------------
def test_security_headers(client):
    r = client.get("/api/health")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["x-frame-options"] == "DENY"
    assert r.headers["referrer-policy"] == "no-referrer"
    assert "default-src 'none'" in r.headers["content-security-policy"]
    post = client.post("/api/compare", json=compare_body())
    assert post.headers["cache-control"] == "no-store"


def test_docs_have_relaxed_csp_and_openapi_is_served(client):
    docs = client.get("/api/docs")
    assert docs.status_code == 200
    assert "cdn.jsdelivr.net" in docs.headers["content-security-policy"]
    assert client.get("/api/openapi.json").json()["info"]["title"] == "BestBill API"


def test_compare_body_limit(client):
    r = client.post(
        "/api/compare",
        content=b"x" * (64 * 1024 + 1),
        headers={"content-type": "application/json"},
    )
    assert r.status_code == 413
    assert r.headers["x-content-type-options"] == "nosniff"


def test_body_limit_streaming_without_content_length(catalog_dir):
    import asyncio

    app = create_app(
        Settings(catalog=CatalogSettings(catalog_path=catalog_dir / "catalog.sqlite"))
    )
    sent: list[dict] = []
    chunks = [
        {"type": "http.request", "body": b"x" * 40_000, "more_body": True},
        {"type": "http.request", "body": b"x" * 40_000, "more_body": False},
    ]

    async def receive():
        return chunks.pop(0)

    async def send(msg):
        sent.append(msg)

    scope = {
        "type": "http",
        "method": "POST",
        "path": "/api/compare",
        "headers": [(b"content-type", b"application/json")],
        "client": ("1.2.3.4", 1),
        "query_string": b"",
        "http_version": "1.1",
        "scheme": "http",
        "server": ("t", 80),
        "root_path": "",
    }
    app.state.provider.load_initial()
    asyncio.run(app(scope, receive, send))
    assert sent[0]["status"] == 413


def test_rate_limit(catalog_dir):
    with make_client(catalog_dir, rate_limit_per_minute=5) as c:
        codes = [c.get("/api/comuni", params={"q": "ro"}).status_code for _ in range(7)]
        assert codes == [200] * 5 + [429] * 2
        limited = c.get("/api/comuni", params={"q": "ro"})
        assert int(limited.headers["retry-after"]) >= 1
        assert limited.headers["x-frame-options"] == "DENY"
        # health checks are exempt
        assert c.get("/api/health").status_code == 200


def test_compare_has_its_own_rate_limit(catalog_dir):
    with make_client(catalog_dir, compare_rate_limit_per_minute=2) as c:
        codes = [
            c.post("/api/compare", json=compare_body()).status_code for _ in range(3)
        ]
        assert codes == [200, 200, 429]
        assert c.get("/api/sample").status_code == 200


def test_rate_limit_ignores_forwarded_for_unless_trusted(catalog_dir):
    def hit(c, ip):
        return c.get("/api/sample", headers={"X-Forwarded-For": f"{ip}, 10.0.0.1"})

    with make_client(catalog_dir, rate_limit_per_minute=2) as c:
        assert [hit(c, f"9.9.9.{i}").status_code for i in range(3)] == [200, 200, 429]
    with make_client(catalog_dir, rate_limit_per_minute=2, trust_proxy=True) as c:
        assert [hit(c, "9.9.9.1").status_code for _ in range(3)] == [200, 200, 429]
        assert hit(c, "9.9.9.2").status_code == 200


def test_rate_limit_key_table_is_bounded():
    from bestbill.api.middleware import RateLimitMiddleware

    async def app(scope, receive, send): ...

    mw = RateLimitMiddleware(app, per_minute=1, path_limits={}, trust_proxy=False)
    mw.MAX_KEYS = 3
    for i in range(10):
        mw._allow(("all", str(i)), 1, 0.0)
    assert len(mw._hits) <= 4


def test_cors_disabled_by_default(client):
    r = client.get("/api/health", headers={"Origin": "https://evil.example"})
    assert "access-control-allow-origin" not in r.headers


def test_cors_allows_only_configured_origins(catalog_dir):
    with make_client(catalog_dir, cors_origins=("https://app.example",)) as c:
        ok = c.get("/api/health", headers={"Origin": "https://app.example"})
        assert ok.headers["access-control-allow-origin"] == "https://app.example"
        other = c.get("/api/health", headers={"Origin": "https://evil.example"})
        assert "access-control-allow-origin" not in other.headers
        pre = c.options(
            "/api/compare",
            headers={
                "Origin": "https://app.example",
                "Access-Control-Request-Method": "POST",
                "Access-Control-Request-Headers": "content-type",
            },
        )
        assert pre.status_code == 200


# -- settings / lifecycle ------------------------------------------------------
def test_settings_from_env():
    s = Settings.from_env(
        {
            "BESTBILL_CATALOG_PATH": "/x/catalog.sqlite",
            "BESTBILL_REPO": "o/r",
            "BESTBILL_RELEASE_TAG": "t",
            "GITHUB_TOKEN": "tok",
            "BESTBILL_CACHE_DIR": "/c",
            "BESTBILL_REFRESH_HOURS": "2",
            "BESTBILL_TRUST_PROXY": "1",
            "BESTBILL_CORS_ORIGINS": "https://a.example, https://b.example,",
        }
    )
    assert s.trust_proxy is True
    assert s.cors_origins == ("https://a.example", "https://b.example")
    assert str(s.catalog.catalog_path) == "/x/catalog.sqlite"
    assert (s.catalog.repo, s.catalog.release_tag) == ("o/r", "t")
    assert s.catalog.github_token == "tok" and s.catalog.refresh_hours == 2
    d = Settings.from_env({})
    assert d.trust_proxy is False and d.cors_origins == ()
    assert d.catalog.catalog_path is None
    assert d.catalog.repo == "leonardoburalli/bestbill"
    assert d.catalog.release_tag == "catalog-latest"
    assert d.catalog.github_token is None and d.catalog.refresh_hours == 6


def test_provider_injection(catalog_dir):
    provider = CatalogProvider(
        CatalogSettings(catalog_path=catalog_dir / "catalog.sqlite")
    )
    with TestClient(create_app(Settings(), provider=provider)) as c:
        assert c.get("/api/health").json()["catalog_loaded"] is True


def test_compare_min_duration_filter_and_fields(client):
    allr = client.post("/api/compare", json=compare_body(top_n=200)).json()
    assert all(
        "duration_months" in x and "duration_open_ended" in x for x in allr["results"]
    )
    assert any(x["duration_open_ended"] for x in allr["results"])
    res = client.post(
        "/api/compare",
        json=compare_body(top_n=200, filters={"min_duration_months": 24}),
    ).json()
    assert res["results"]
    assert res["total_matching"] < res["total_eligible"]
    for x in res["results"]:
        assert not x["duration_open_ended"]
        assert x["duration_months"] >= 24


def test_compare_min_duration_validation(client):
    for bad_value in (0, 121):
        r = client.post(
            "/api/compare",
            json=compare_body(filters={"min_duration_months": bad_value}),
        )
        assert r.status_code == 422


def test_compare_exact_and_over_duration(client):
    exact = client.post(
        "/api/compare",
        json=compare_body(
            top_n=200, filters={"min_duration_months": 12, "max_duration_months": 12}
        ),
    ).json()
    assert exact["results"]
    assert all(
        x["duration_months"] == 12 and not x["duration_open_ended"]
        for x in exact["results"]
    )
    over = client.post(
        "/api/compare",
        json=compare_body(top_n=200, filters={"min_duration_months": 13}),
    ).json()
    for x in over["results"]:
        assert not x["duration_open_ended"]
        assert x["duration_months"] > 12


def test_compare_duration_bounds_validation(client):
    r = client.post(
        "/api/compare",
        json=compare_body(
            filters={"min_duration_months": 24, "max_duration_months": 12}
        ),
    )
    assert r.status_code == 422
    r = client.post(
        "/api/compare", json=compare_body(filters={"max_duration_months": 0})
    )
    assert r.status_code == 422


def test_offers_list_has_duration_fields(client):
    items = client.get("/api/offers?limit=5").json()["items"]
    assert all("duration_months" in i and "duration_open_ended" in i for i in items)
