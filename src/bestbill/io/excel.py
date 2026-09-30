"""Read the legacy Excel format (``Tariffario`` + ``Storico_<Location>``
sheets) into core domain objects, using openpyxl directly (no pandas).
"""

from __future__ import annotations

import pathlib
from datetime import date, datetime
from typing import IO

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from bestbill.core.models import (
    BandStructure,
    ConsumptionProfile,
    MonthlyConsumption,
    Offer,
    OfferSource,
    PriceType,
    PunSeries,
)

MAX_SHEETS = 50
MAX_ROWS = 1000

STORICO_PREFIX = "Storico_"
TARIFFARIO_SHEET = "Tariffario"

_TARIFFARIO_COLUMNS = (
    "Tariffa",
    "Prezzo fisso [€/kWh]",
    "Alpha [€/kWh]",
    "CCV [€]",
)
_STORICO_COLUMNS = (
    "Mese",
    "Consumo [kWh]",
    "PUN mensile [€/kWh]",
    "FC PUN mensile [€/kWh]",
)


class ExcelFormatError(ValueError):
    """The workbook doesn't match the expected legacy format."""


def _open_workbook(path: str | pathlib.Path | IO[bytes]) -> openpyxl.workbook.Workbook:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError as exc:
        raise ExcelFormatError(f"file not found: {path}") from exc
    except OSError as exc:
        raise ExcelFormatError(f"could not open {path}: {exc}") from exc
    if len(wb.sheetnames) > MAX_SHEETS:
        raise ExcelFormatError(
            f"workbook has {len(wb.sheetnames)} sheets, more than the {MAX_SHEETS} cap"
        )
    return wb


def _header_index(ws: Worksheet, expected_columns: tuple[str, ...]) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise ExcelFormatError(f"sheet {ws.title!r} is empty")
    index = {str(name): i for i, name in enumerate(header_row) if name is not None}
    missing = [c for c in expected_columns if c not in index]
    if missing:
        raise ExcelFormatError(
            f"sheet {ws.title!r} is missing column(s) {missing}; found {list(index)}"
        )
    return index


def _iter_data_rows(ws: Worksheet) -> list[tuple[object, ...]]:
    rows = list(ws.iter_rows(min_row=2, max_row=MAX_ROWS + 1, values_only=True))
    if len(rows) > MAX_ROWS:
        raise ExcelFormatError(
            f"sheet {ws.title!r} has more than the {MAX_ROWS}-row cap"
        )
    return [row for row in rows if any(v is not None for v in row)]


def list_locations(path: str | pathlib.Path | IO[bytes]) -> list[str]:
    """Return the locations available in the workbook (from ``Storico_*``
    sheet names), preserving sheet order.
    """
    wb = _open_workbook(path)
    return [
        name[len(STORICO_PREFIX) :]
        for name in wb.sheetnames
        if name.startswith(STORICO_PREFIX)
    ]


def _to_date(value: object, *, context: str) -> date:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    raise ExcelFormatError(f"{context}: expected a date, got {value!r}")


def _to_float(value: object, *, context: str) -> float:
    if isinstance(value, int | float):
        return float(value)
    raise ExcelFormatError(f"{context}: expected a number, got {value!r}")


def read_custom_offers(path: str | pathlib.Path) -> list[Offer]:
    """Read the ``Tariffario`` sheet into custom (mono-band) offers.

    ``Prezzo fisso [€/kWh] == 0`` means a variable offer priced as
    PUN + Alpha; otherwise it's a fixed offer priced at Prezzo fisso + Alpha.
    """
    wb = _open_workbook(path)
    if TARIFFARIO_SHEET not in wb.sheetnames:
        raise ExcelFormatError(
            f"sheet {TARIFFARIO_SHEET!r} not found; available: {wb.sheetnames}"
        )
    ws = wb[TARIFFARIO_SHEET]
    cols = _header_index(ws, _TARIFFARIO_COLUMNS)

    offers: list[Offer] = []
    for row in _iter_data_rows(ws):
        name = row[cols["Tariffa"]]
        if name is None:
            continue
        name = str(name)
        ctx = f"{TARIFFARIO_SHEET!r} row for tariff {name!r}"
        prezzo_fisso = _to_float(row[cols["Prezzo fisso [€/kWh]"]], context=ctx)
        alpha = _to_float(row[cols["Alpha [€/kWh]"]], context=ctx)
        ccv = _to_float(row[cols["CCV [€]"]], context=ctx)

        offer_id = f"custom:{name}"
        if prezzo_fisso == 0:
            offers.append(
                Offer(
                    id=offer_id,
                    supplier=name,
                    name=name,
                    source=OfferSource.CUSTOM,
                    price_type=PriceType.VARIABLE,
                    band_structure=BandStructure.MONO,
                    energy_price_eur_kwh={},
                    spread_eur_kwh={"mono": alpha},
                    fixed_fee_eur_year=ccv,
                )
            )
        else:
            offers.append(
                Offer(
                    id=offer_id,
                    supplier=name,
                    name=name,
                    source=OfferSource.CUSTOM,
                    price_type=PriceType.FIXED,
                    band_structure=BandStructure.MONO,
                    energy_price_eur_kwh={"mono": prezzo_fisso},
                    spread_eur_kwh={"mono": alpha},
                    fixed_fee_eur_year=ccv,
                )
            )
    return offers


def read_location_data(
    path: str | pathlib.Path | IO[bytes], location: str
) -> tuple[ConsumptionProfile, PunSeries, PunSeries]:
    """Read the ``Storico_<location>`` sheet: 12 months of consumption, the
    actual PUN series, and the forecast PUN series.
    """
    wb = _open_workbook(path)
    sheet_name = f"{STORICO_PREFIX}{location}"
    if sheet_name not in wb.sheetnames:
        available = list_locations(path)
        raise ExcelFormatError(
            f"sheet {sheet_name!r} not found; available locations: {available}"
        )
    ws = wb[sheet_name]
    cols = _header_index(ws, _STORICO_COLUMNS)

    months: list[MonthlyConsumption] = []
    actual_pun: dict[date, float] = {}
    forecast_pun: dict[date, float] = {}
    for row in _iter_data_rows(ws):
        ctx = f"{sheet_name!r} row"
        month = _to_date(row[cols["Mese"]], context=ctx)
        kwh = _to_float(row[cols["Consumo [kWh]"]], context=ctx)
        pun = _to_float(row[cols["PUN mensile [€/kWh]"]], context=ctx)
        fc_pun = _to_float(row[cols["FC PUN mensile [€/kWh]"]], context=ctx)

        months.append(MonthlyConsumption(month=month, kwh=kwh))
        actual_pun[month] = pun
        forecast_pun[month] = fc_pun

    profile = ConsumptionProfile(months=months)
    return profile, PunSeries(values=actual_pun), PunSeries(values=forecast_pun)
